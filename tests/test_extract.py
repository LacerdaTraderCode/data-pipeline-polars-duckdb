import json

import polars as pl
from openpyxl import Workbook

from pipeline.extract import read_csv, read_excel, read_json, read_parquet, scan_parquet


def test_read_csv(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("id,name\n1,a\n2,b\n", encoding="utf-8")

    df = read_csv(str(path))

    assert df.to_dicts() == [{"id": 1, "name": "a"}, {"id": 2, "name": "b"}]


def test_read_csv_forwards_options(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("id;name\n1;a\n", encoding="utf-8")

    df = read_csv(str(path), separator=";")

    assert df.columns == ["id", "name"]


def test_read_json(tmp_path):
    path = tmp_path / "data.json"
    path.write_text(json.dumps([{"id": 1, "name": "a"}, {"id": 2, "name": "b"}]), encoding="utf-8")

    df = read_json(str(path))

    assert df.to_dicts() == [{"id": 1, "name": "a"}, {"id": 2, "name": "b"}]


def test_read_excel_defaults_to_first_sheet(tmp_path):
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["id", "name"])
    workbook.active.append([1, "a"])
    workbook.active.append([2, "b"])
    workbook.save(path)

    df = read_excel(str(path))

    assert df.columns == ["id", "name"]
    assert df["name"].to_list() == ["a", "b"]


def test_read_excel_selects_named_sheet(tmp_path):
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.title = "first"
    workbook.active.append(["x"])
    workbook.active.append([1])
    second = workbook.create_sheet("second")
    second.append(["y"])
    second.append([2])
    workbook.save(path)

    df = read_excel(str(path), sheet="second")

    assert df.columns == ["y"]


def test_read_parquet_roundtrip(tmp_path, sales):
    path = tmp_path / "sales.parquet"
    sales.write_parquet(path)

    assert read_parquet(str(path)).equals(sales)


def test_scan_parquet_is_lazy(sales_parquet, sales):
    lazy = scan_parquet(sales_parquet)

    assert isinstance(lazy, pl.LazyFrame)
    assert lazy.collect().equals(sales)
